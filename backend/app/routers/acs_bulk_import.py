"""Feature 11: Bulk Import Persons + Cards from CSV/Excel."""

import csv
import io
import json
import logging
from typing import Optional

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from sqlalchemy.orm import Session

from app.auth import require_admin, require_operator_or_admin
from app.db import get_db
from app.models import AccessControlDevice, CardIssue, Person, User
from app.utils import utc_now

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/api/acs/bulk-import", tags=["acs_bulk_import"])


def _clean(v) -> Optional[str]:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _parse_csv_rows(content: bytes) -> list[dict]:
    text = content.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    return [row for row in reader]


def _parse_xlsx_rows(content: bytes) -> list[dict]:
    try:
        import openpyxl
    except ImportError:
        raise HTTPException(422, "openpyxl is required for Excel import; use CSV instead")
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
    ws = wb.active
    headers = [str(cell.value or "").strip() for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        rows.append(dict(zip(headers, row)))
    return rows


@router.get("/template/csv")
def download_template(current_user: User = Depends(require_operator_or_admin)):
    """Return a CSV template for bulk import."""
    from fastapi.responses import StreamingResponse
    header = "full_name,employee_no,department,email,phone,position,card_number,valid_from,valid_to\n"
    example = "John Doe,EMP001,HR,john@company.com,+37491000000,Manager,1234567890,2025-01-01,2025-12-31\n"
    return StreamingResponse(
        iter([header, example]),
        media_type="text/csv",
        headers={"Content-Disposition": 'attachment; filename="acs_bulk_import_template.csv"'},
    )


@router.post("/preview")
async def preview_import(
    file: UploadFile = File(...),
    device_id: Optional[int] = None,
    current_user: User = Depends(require_operator_or_admin),
    db: Session = Depends(get_db),
):
    """Parse the file and return a preview without writing to DB."""
    content = await file.read()
    if file.filename.endswith(".xlsx") or file.filename.endswith(".xls"):
        rows = _parse_xlsx_rows(content)
    else:
        rows = _parse_csv_rows(content)

    preview = []
    for i, row in enumerate(rows[:200]):
        employee_no = _clean(row.get("employee_no"))
        full_name   = _clean(row.get("full_name"))
        card_number = _clean(row.get("card_number"))
        exists = False
        if employee_no:
            exists = bool(db.query(Person).filter(Person.employee_no == employee_no).first())
        preview.append({
            "row": i + 2,
            "full_name": full_name,
            "employee_no": employee_no,
            "department": _clean(row.get("department")),
            "email": _clean(row.get("email")),
            "phone": _clean(row.get("phone")),
            "position": _clean(row.get("position")),
            "card_number": card_number,
            "valid_from": _clean(row.get("valid_from")),
            "valid_to": _clean(row.get("valid_to")),
            "action": "update" if exists else "create",
        })

    return {
        "total_rows": len(rows),
        "preview_rows": len(preview),
        "columns_found": list(rows[0].keys()) if rows else [],
        "preview": preview,
    }


@router.post("/execute")
async def execute_import(
    file: UploadFile = File(...),
    device_id: Optional[int] = None,
    update_existing: bool = True,
    current_user: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    """Execute the bulk import — create/update persons and issue cards."""
    content = await file.read()
    if file.filename.endswith(".xlsx") or file.filename.endswith(".xls"):
        rows = _parse_xlsx_rows(content)
    else:
        rows = _parse_csv_rows(content)

    device = db.get(AccessControlDevice, device_id) if device_id else None

    created = updated = skipped = errors = 0
    error_rows: list = []

    for i, row in enumerate(rows):
        try:
            full_name   = _clean(row.get("full_name"))
            employee_no = _clean(row.get("employee_no"))
            card_number = _clean(row.get("card_number"))

            if not full_name and not employee_no:
                skipped += 1
                continue

            # Find or create person
            person = None
            if employee_no:
                person = db.query(Person).filter(Person.employee_no == employee_no).first()

            if person and not update_existing:
                skipped += 1
                continue

            if not person:
                person = Person(
                    full_name=full_name,
                    employee_no=employee_no,
                    person_type="employee",
                    source="acs",
                )
                db.add(person)
                db.flush()
                created += 1
            else:
                if full_name:
                    person.full_name = full_name
                updated += 1

            # Update HR fields
            for field in ("department", "email", "phone", "position"):
                val = _clean(row.get(field))
                if val:
                    setattr(person, field, val)

            # Issue card if provided and device specified
            if card_number and device:
                existing_card = db.query(CardIssue).filter(
                    CardIssue.person_id == person.id,
                    CardIssue.device_id == device.id,
                    CardIssue.card_number == card_number,
                    CardIssue.status == "active",
                ).first()
                if not existing_card:
                    from datetime import datetime
                    valid_from_str = _clean(row.get("valid_from"))
                    valid_to_str   = _clean(row.get("valid_to"))
                    card = CardIssue(
                        person_id=person.id,
                        device_id=device.id,
                        card_number=card_number,
                        card_type="normal",
                        valid_from=datetime.fromisoformat(valid_from_str) if valid_from_str else None,
                        valid_to=datetime.fromisoformat(valid_to_str) if valid_to_str else None,
                        status="active",
                    )
                    db.add(card)

        except Exception as exc:
            errors += 1
            error_rows.append({"row": i + 2, "error": str(exc)})
            logger.warning("Bulk import row %d error: %s", i + 2, exc)

    db.commit()

    return {
        "ok": True,
        "total_rows": len(rows),
        "created": created,
        "updated": updated,
        "skipped": skipped,
        "errors": errors,
        "error_rows": error_rows[:50],
    }
