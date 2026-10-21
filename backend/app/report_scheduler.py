"""Generation, storage, delivery, and scheduled execution of Face Analytics reports."""

import csv
import io
import json
import logging
import os
import smtplib
import threading
import time
from datetime import datetime, timedelta, timezone
from email import encoders
from email.mime.base import MIMEBase
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from pathlib import Path
from typing import Any, Optional

from sqlalchemy import func
from sqlalchemy.orm import joinedload

logger = logging.getLogger(__name__)

_scheduler_started = False
_scheduler_lock = threading.Lock()

REPORT_TYPES = {
    "attendance": {"title": "Attendance and visit timeline", "description": "Recognized visits, directions, and attendance evidence"},
    "events": {"title": "Face events", "description": "All face recognition events in the selected period"},
    "unknowns": {"title": "Unknown persons", "description": "Unidentified detections for review"},
    "persons": {"title": "Persons database", "description": "Registered persons and recognition activity"},
    "frequency": {"title": "Person frequency", "description": "Most frequently detected persons"},
    "camera_activity": {"title": "Camera activity", "description": "Detection volume by camera"},
    "quality": {"title": "Recognition quality", "description": "Face quality and low-quality detections"},
    "audit": {"title": "Audit log", "description": "Administrator and operator activity"},
}
VALID_FORMATS = {"excel", "pdf", "csv"}


def _matches_cron(cron_expr: str, dt: datetime) -> bool:
    """Small cron checker for the platform's five-field schedules."""
    try:
        minute, hour, dom, month, dow = cron_expr.strip().split()

        def match(field: str, value: int) -> bool:
            if field == "*":
                return True
            if field.isdigit():
                return int(field) == value
            if "," in field:
                return value in [int(item) for item in field.split(",") if item.isdigit()]
            if "/" in field:
                start, step = field.split("/", 1)
                step_value = int(step)
                start_value = 0 if start == "*" else int(start)
                return step_value > 0 and (value - start_value) % step_value == 0
            return False

        return all((match(minute, dt.minute), match(hour, dt.hour), match(dom, dt.day), match(month, dt.month), match(dow, dt.weekday())))
    except Exception:
        return False


def normalize_settings(value: Any) -> dict:
    default = {"period": "last_7_days", "filters": {}, "language": "en", "delivery": {"email": True, "save_local": True}}
    if isinstance(value, str):
        try:
            value = json.loads(value or "{}")
        except (TypeError, ValueError):
            value = {}
    if not isinstance(value, dict):
        return default
    delivery = value.get("delivery") if isinstance(value.get("delivery"), dict) else {}
    filters = value.get("filters") if isinstance(value.get("filters"), dict) else {}
    return {
        "period": value.get("period") or default["period"],
        "filters": filters,
        "language": value.get("language") if value.get("language") in {"en", "ru", "hy"} else default["language"],
        "delivery": {"email": bool(delivery.get("email", True)), "save_local": bool(delivery.get("save_local", True))},
    }


def _to_naive(value: Optional[datetime]) -> Optional[datetime]:
    if value is None:
        return None
    return value.replace(tzinfo=None) if value.tzinfo else value


def _parse_datetime(value: Any, end_of_day: bool = False) -> Optional[datetime]:
    if not value:
        return None
    try:
        raw = str(value).strip().replace("Z", "+00:00")
        parsed = datetime.fromisoformat(raw)
        if len(raw) == 10 and end_of_day:
            parsed = parsed.replace(hour=23, minute=59, second=59, microsecond=999999)
        return _to_naive(parsed)
    except ValueError:
        return None


def resolve_filters(value: Optional[dict] = None) -> dict:
    filters = dict(value or {})
    now = datetime.now(timezone.utc)
    if not filters.get("date_from") and not filters.get("date_to"):
        period = filters.get("period", "last_7_days")
        if period == "today":
            start = now.replace(hour=0, minute=0, second=0, microsecond=0)
        elif period == "last_30_days":
            start = now - timedelta(days=30)
        elif period == "this_month":
            start = now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
        else:
            start = now - timedelta(days=7)
        filters["date_from"] = start.date().isoformat()
        filters["date_to"] = now.date().isoformat()
    return filters


def _selected_person_ids(filters: dict) -> list[int]:
    """Return only valid person IDs from the report filter payload."""
    raw_values = filters.get("person_ids") or ([filters.get("person_id")] if filters.get("person_id") else [])
    if not isinstance(raw_values, (list, tuple, set)):
        raw_values = [raw_values]
    result: list[int] = []
    for value in raw_values:
        try:
            person_id = int(value)
        except (TypeError, ValueError):
            continue
        if person_id > 0 and person_id not in result:
            result.append(person_id)
    return result


def _face_events_query(db, filters: dict):
    from app.models import FaceEvent, Person

    query = db.query(FaceEvent).options(joinedload(FaceEvent.person)).filter(FaceEvent.source == "face_analytics")
    date_from = _parse_datetime(filters.get("date_from"))
    date_to = _parse_datetime(filters.get("date_to"), end_of_day=True)
    if date_from:
        query = query.filter(FaceEvent.created_at >= date_from)
    if date_to:
        query = query.filter(FaceEvent.created_at <= date_to)
    if filters.get("camera_id"):
        try:
            query = query.filter(FaceEvent.camera_id == int(filters["camera_id"]))
        except (TypeError, ValueError):
            pass
    if filters.get("camera_name"):
        query = query.filter(FaceEvent.camera_name.ilike(f"%{str(filters['camera_name']).strip()}%"))
    if filters.get("department"):
        query = query.filter(FaceEvent.person.has(Person.department.ilike(f"%{str(filters['department']).strip()}%")))
    person_ids = _selected_person_ids(filters)
    if person_ids:
        query = query.filter(FaceEvent.person_id.in_(person_ids))
    if filters.get("direction") in {"in", "out", "both"}:
        query = query.filter(FaceEvent.direction == filters["direction"])
    return query


def collect_report_data(report_type: str, db, filters: Optional[dict] = None) -> tuple[list[str], list[list[Any]], dict]:
    """Return columns, rows, and dashboard-friendly summary for a report type."""
    from app.models import AuditLog, FaceEvent, Person

    filters = resolve_filters(filters)
    if report_type not in REPORT_TYPES:
        raise ValueError("Unsupported report type")

    if report_type in {"attendance", "events", "unknowns", "quality"}:
        query = _face_events_query(db, filters)
        if report_type == "unknowns":
            query = query.filter(FaceEvent.person_id.is_(None))
        events = query.order_by(FaceEvent.created_at.desc()).limit(10000).all()
        if report_type == "quality":
            headers = ["Event ID", "Person", "Camera", "Time", "Confidence", "Face Quality", "Low Quality", "Reason"]
            rows = [[event.id, event.person.full_name if event.person else "Unknown", event.camera_name or "-", str(event.created_at), event.confidence, event.face_quality, "Yes" if event.is_low_quality else "No", event.quality_reason or "-"] for event in events]
            quality_values = [float(event.face_quality) for event in events if event.face_quality is not None]
            return headers, rows, {"events": len(rows), "average_quality": round(sum(quality_values) / len(quality_values), 2) if quality_values else None, "low_quality": sum(1 for event in events if event.is_low_quality)}
        headers = ["Event ID", "Person", "Person Type", "Camera", "Direction", "Time", "Confidence", "Face Quality"]
        rows = [[event.id, event.person.full_name if event.person and event.person.full_name else "Unknown", event.person.person_type if event.person else "unknown", event.camera_name or "-", event.direction or "-", str(event.created_at), event.confidence, event.face_quality] for event in events]
        return headers, rows, {"events": len(rows), "recognized": sum(1 for event in events if event.person_id), "unknown": sum(1 for event in events if not event.person_id)}

    if report_type == "persons":
        query = db.query(Person).filter(Person.full_name.isnot(None))
        person_ids = _selected_person_ids(filters)
        if person_ids:
            query = query.filter(Person.id.in_(person_ids))
        if filters.get("department"):
            query = query.filter(Person.department.ilike(f"%{str(filters['department']).strip()}%"))
        if filters.get("person_type"):
            query = query.filter(Person.person_type == filters["person_type"])
        persons = query.order_by(Person.last_seen.desc()).limit(10000).all()
        headers = ["ID", "Full Name", "Type", "Department", "Position", "First Seen", "Last Seen", "Seen Count"]
        rows = [[person.id, person.full_name, person.person_type, person.department or "-", person.position or "-", str(person.first_seen), str(person.last_seen), person.seen_count] for person in persons]
        return headers, rows, {"persons": len(rows), "employees": sum(1 for person in persons if person.person_type == "employee")}

    if report_type == "frequency":
        events = _face_events_query(db, filters).filter(FaceEvent.person_id.isnot(None)).all()
        counts: dict[int, dict] = {}
        for event in events:
            person = event.person
            if not person:
                continue
            item = counts.setdefault(person.id, {"name": person.full_name or f"Person {person.id}", "department": person.department or "-", "count": 0, "first": event.created_at, "last": event.created_at})
            item["count"] += 1
            item["first"] = min(item["first"], event.created_at)
            item["last"] = max(item["last"], event.created_at)
        values = sorted(counts.values(), key=lambda item: item["count"], reverse=True)
        headers = ["Person", "Department", "Detections", "First Detection", "Last Detection"]
        rows = [[item["name"], item["department"], item["count"], str(item["first"]), str(item["last"])] for item in values]
        return headers, rows, {"unique_persons": len(rows), "detections": sum(item["count"] for item in values)}

    if report_type == "camera_activity":
        events = _face_events_query(db, filters).all()
        counts: dict[str, dict] = {}
        for event in events:
            name = event.camera_name or "Unnamed camera"
            item = counts.setdefault(name, {"events": 0, "recognized": 0, "unknown": 0, "first": event.created_at, "last": event.created_at})
            item["events"] += 1
            item["recognized"] += 1 if event.person_id else 0
            item["unknown"] += 0 if event.person_id else 1
            item["first"] = min(item["first"], event.created_at)
            item["last"] = max(item["last"], event.created_at)
        headers = ["Camera", "Events", "Recognized", "Unknown", "First Event", "Last Event"]
        rows = [[name, item["events"], item["recognized"], item["unknown"], str(item["first"]), str(item["last"])] for name, item in sorted(counts.items(), key=lambda pair: pair[1]["events"], reverse=True)]
        return headers, rows, {"cameras": len(rows), "events": sum(item["events"] for item in counts.values())}

    query = db.query(AuditLog)
    date_from = _parse_datetime(filters.get("date_from"))
    date_to = _parse_datetime(filters.get("date_to"), end_of_day=True)
    if date_from:
        query = query.filter(AuditLog.created_at >= date_from)
    if date_to:
        query = query.filter(AuditLog.created_at <= date_to)
    logs = query.order_by(AuditLog.created_at.desc()).limit(10000).all()
    headers = ["ID", "Username", "Action", "Target", "Details", "IP", "Time"]
    rows = [[log.id, log.username or "-", log.action, log.target_type or "-", log.details or "-", log.ip_address or "-", str(log.created_at)] for log in logs]
    return headers, rows, {"actions": len(rows), "operators": len({log.username for log in logs if log.username})}


def dashboard_data(db, filters: Optional[dict] = None) -> dict:
    from app.models import FaceEvent

    filters = resolve_filters(filters)
    events = _face_events_query(db, filters).order_by(FaceEvent.created_at.asc()).all()
    known = [event for event in events if event.person_id]
    quality = [float(event.face_quality) for event in events if event.face_quality is not None]
    daily: dict[str, dict] = {}
    cameras: dict[str, int] = {}
    for event in events:
        day = event.created_at.date().isoformat() if event.created_at else "-"
        bucket = daily.setdefault(day, {"date": day, "events": 0, "recognized": 0, "unknown": 0})
        bucket["events"] += 1
        bucket["recognized"] += 1 if event.person_id else 0
        bucket["unknown"] += 0 if event.person_id else 1
        camera = event.camera_name or "Unnamed camera"
        cameras[camera] = cameras.get(camera, 0) + 1
    return {
        "filters": filters,
        "summary": {"events": len(events), "recognized": len(known), "unknown": len(events) - len(known), "unique_persons": len({event.person_id for event in known}), "average_quality": round(sum(quality) / len(quality), 1) if quality else None},
        "daily": list(daily.values())[-31:],
        "cameras": [{"name": name, "events": count} for name, count in sorted(cameras.items(), key=lambda item: item[1], reverse=True)[:10]],
    }


def _generate_report(report_type: str, fmt: str, db, filters: Optional[dict] = None) -> tuple[bytes, str, str, int]:
    if report_type not in REPORT_TYPES:
        raise ValueError("Unsupported report type")
    if fmt not in VALID_FORMATS:
        raise ValueError("Unsupported report format")
    headers, rows, summary = collect_report_data(report_type, db, filters)
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    safe_type = report_type.replace("_", "-")

    if fmt == "excel":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = REPORT_TYPES[report_type]["title"][:31]
        sheet.append([REPORT_TYPES[report_type]["title"]])
        sheet.append(["Generated (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")])
        for key, value in summary.items():
            sheet.append([str(key).replace("_", " ").title(), value])
        sheet.append([])
        sheet.append(headers)
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="5146ED")
        for row in rows:
            sheet.append(row)
        sheet.freeze_panes = "A7"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(44, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue(), f"{safe_type}_report_{timestamp}.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", len(rows)

    if fmt == "csv":
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow([REPORT_TYPES[report_type]["title"]])
        writer.writerow(["Generated (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")])
        for key, value in summary.items():
            writer.writerow([str(key).replace("_", " ").title(), value])
        writer.writerow([])
        writer.writerow(headers)
        writer.writerows(rows)
        return output.getvalue().encode("utf-8-sig"), f"{safe_type}_report_{timestamp}.csv", "text/csv; charset=utf-8", len(rows)

    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError as error:
        raise RuntimeError("PDF report export is unavailable: reportlab is not installed") from error

    output = io.BytesIO()
    document = SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=12 * mm)
    styles = getSampleStyleSheet()
    body = [Paragraph(REPORT_TYPES[report_type]["title"], styles["Title"]), Paragraph(f"Generated (UTC): {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S')}", styles["Normal"]), Spacer(1, 5 * mm)]
    if summary:
        body.append(Table([[str(key).replace("_", " ").title(), str(value)] for key, value in summary.items()], colWidths=[55 * mm, 45 * mm], style=TableStyle([("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#DCE4F0")), ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F0EFFF"))])))
        body.append(Spacer(1, 5 * mm))
    pdf_rows = [headers] + [[str(value if value is not None else "")[:48] for value in row] for row in rows[:500]]
    widths = [((273 * mm) / max(len(headers), 1))] * len(headers)
    table = Table(pdf_rows, colWidths=widths, repeatRows=1)
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#5146ED")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 6.5), ("GRID", (0, 0), (-1, -1), 0.2, colors.HexColor("#DCE4F0")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    body.append(table)
    document.build(body)
    return output.getvalue(), f"{safe_type}_report_{timestamp}.pdf", "application/pdf", len(rows)


def _write_report_file(run_id: int, filename: str, data: bytes) -> tuple[str, str]:
    directory = Path(os.getenv("REPORT_EXPORT_DIR", "exports/reports"))
    directory.mkdir(parents=True, exist_ok=True)
    stored_name = f"{run_id}_{filename}"
    path = directory / stored_name
    path.write_bytes(data)
    return stored_name, str(path)


def generate_report_run(db, report_type: str, fmt: str, report_name: str, filters: Optional[dict] = None, scheduled_report_id: Optional[int] = None, initiated_by: Optional[str] = None) -> tuple[object, bytes, str, str]:
    from app.models import ReportRun
    from app.utils import utc_now

    normalized = resolve_filters(filters)
    run = ReportRun(scheduled_report_id=scheduled_report_id, report_name=report_name, report_type=report_type, format=fmt, status="running", parameters=json.dumps(normalized), initiated_by=initiated_by, started_at=utc_now())
    db.add(run)
    db.flush()
    try:
        data, filename, mime, row_count = _generate_report(report_type, fmt, db, normalized)
        stored_name, stored_path = _write_report_file(run.id, filename, data)
        run.file_name = stored_name
        run.file_path = stored_path
        run.row_count = row_count
        run.status = "generated"
        run.completed_at = utc_now()
        db.commit()
        db.refresh(run)
        return run, data, filename, mime
    except Exception as error:
        db.rollback()
        failed = ReportRun(
            scheduled_report_id=scheduled_report_id,
            report_name=report_name,
            report_type=report_type,
            format=fmt,
            status="failed",
            parameters=json.dumps(normalized),
            initiated_by=initiated_by,
            message=str(error)[:1000],
            started_at=utc_now(),
            completed_at=utc_now(),
        )
        db.add(failed)
        db.commit()
        raise


def _send_report_email(recipients: list[str], report_name: str, data: bytes, filename: str, mime_type: str) -> bool:
    smtp_host = os.getenv("SMTP_HOST", "")
    smtp_port = int(os.getenv("SMTP_PORT", "587"))
    smtp_user = os.getenv("SMTP_USER", "")
    smtp_pass = os.getenv("SMTP_PASS", "")
    smtp_from = os.getenv("SMTP_FROM", smtp_user)
    use_tls = os.getenv("SMTP_TLS", "1") == "1"
    if not smtp_host:
        logger.warning("SMTP_HOST is not configured; report was saved locally only")
        return False
    try:
        message = MIMEMultipart()
        message["Subject"] = f"[Face Analytics Pro] {report_name}"
        message["From"] = smtp_from
        message["To"] = ", ".join(recipients)
        message.attach(MIMEText("The requested Face Analytics Pro report is attached.", "plain"))
        part = MIMEBase("application", "octet-stream")
        part.set_payload(data)
        encoders.encode_base64(part)
        part.add_header("Content-Disposition", f'attachment; filename="{filename}"')
        message.attach(part)
        server = smtplib.SMTP(smtp_host, smtp_port, timeout=15)
        if use_tls:
            server.ehlo()
            server.starttls()
        if smtp_user and smtp_pass:
            server.login(smtp_user, smtp_pass)
        server.sendmail(smtp_from, recipients, message.as_string())
        server.quit()
        return True
    except Exception as error:
        logger.error("Could not deliver report by email: %s", error)
        return False


def execute_scheduled_report(report_id: int) -> None:
    from app.db import SessionLocal
    from app.models import ScheduledReport
    from app.utils import utc_now

    db = SessionLocal()
    try:
        report = db.get(ScheduledReport, report_id)
        if not report:
            return
        settings = normalize_settings(report.settings)
        run, data, filename, mime = generate_report_run(db, report.report_type, report.format, report.name, settings.get("filters"), report.id, "scheduler")
        recipients = json.loads(report.recipients or "[]")
        should_email = bool(settings["delivery"].get("email", True) and recipients)
        delivered = _send_report_email(recipients, report.name, data, filename, mime) if should_email else False
        report = db.get(ScheduledReport, report_id)
        run = db.get(type(run), run.id)
        run.delivery_status = "emailed" if delivered else ("not_configured" if not should_email else "failed")
        run.status = "delivered" if delivered else ("generated" if not should_email else "delivery_failed")
        report.last_run_status = run.status
        report.last_run_message = "Email delivered" if delivered else ("Saved to report history" if not should_email else "Email delivery failed; file is saved in history")
        if delivered:
            report.last_sent_at = utc_now()
        db.commit()
    except Exception as error:
        logger.exception("Scheduled report %s failed", report_id)
        try:
            report = db.get(ScheduledReport, report_id)
            if report:
                report.last_run_status = "failed"
                report.last_run_message = str(error)[:1000]
                db.commit()
        except Exception:
            db.rollback()
    finally:
        db.close()


def _run_due_reports(db) -> None:
    from app.models import ScheduledReport

    now = datetime.now(timezone.utc)
    reports = db.query(ScheduledReport).filter(ScheduledReport.is_active == 1).all()
    for report in reports:
        if _matches_cron(report.schedule_cron or "0 8 * * 1", now):
            threading.Thread(target=execute_scheduled_report, args=(report.id,), daemon=True).start()


def start_report_scheduler() -> None:
    global _scheduler_started
    with _scheduler_lock:
        if _scheduler_started:
            return
        _scheduler_started = True

    def loop():
        from app.db import SessionLocal
        logger.info("Report scheduler started")
        last_minute = -1
        while True:
            try:
                now = datetime.now(timezone.utc)
                if now.minute != last_minute:
                    last_minute = now.minute
                    from app.distributed import coordinator
                    with coordinator.lock(f"reports:{now.strftime('%Y%m%d%H%M')}", ttl_seconds=90) as acquired:
                        if acquired:
                            db = SessionLocal()
                            try:
                                _run_due_reports(db)
                            finally:
                                db.close()
            except Exception as error:
                logger.exception("Report scheduler loop error: %s", error)
            time.sleep(30)

    threading.Thread(target=loop, daemon=True).start()
