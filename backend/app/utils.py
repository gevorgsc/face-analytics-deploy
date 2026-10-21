"""Shared utilities used across routers - extracted from main.py to avoid duplication."""

import os
import json
import logging
import asyncio
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone
from typing import Optional
from zoneinfo import ZoneInfo

import cv2
from fastapi import HTTPException, UploadFile
from sqlalchemy.orm import Session

from app.face_engine import extract_faces, compare_embeddings


def _create_media_url(path):
    """Lazy import to avoid circular deps and startup issues."""
    from app.media_security import create_media_url
    return create_media_url(path)

logger = logging.getLogger(__name__)

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MEDIA_DIR = os.path.join(BASE_DIR, "media")
EXPORT_DIR = os.path.join(BASE_DIR, "exports")

os.makedirs(MEDIA_DIR, exist_ok=True)
os.makedirs(EXPORT_DIR, exist_ok=True)

try:
    YEREVAN_TZ = ZoneInfo("Asia/Yerevan")
except Exception:
    YEREVAN_TZ = timezone(timedelta(hours=4))

MAX_UPLOAD_BYTES = int(os.getenv("MAX_UPLOAD_MB", "10")) * 1024 * 1024
MATCH_THRESHOLD = 0.45
UNKNOWN_GROUP_THRESHOLD = 0.38
MIN_FACE_DET_SCORE = 0.35
MIN_FACE_SIZE = 35
MIN_FACE_QUALITY_SCORE = int(os.getenv("FACE_QUALITY_THRESHOLD", "30"))


def get_face_quality_threshold(db=None) -> int:
    """Return the effective face quality threshold, checking AppSetting if db provided."""
    if db is not None:
        try:
            from app.models import AppSetting
            row = db.get(AppSetting, "face_quality_threshold")
            if row and row.value:
                return int(float(row.value))
        except Exception:
            pass
    return MIN_FACE_QUALITY_SCORE

_face_executor = ThreadPoolExecutor(max_workers=int(os.getenv("FACE_WORKERS", "2")))


def utc_now() -> datetime:
    return datetime.now(timezone.utc)


def yerevan_day_start_utc() -> datetime:
    local_now = utc_now().astimezone(YEREVAN_TZ)
    local_start = datetime(local_now.year, local_now.month, local_now.day, tzinfo=YEREVAN_TZ)
    return local_start.astimezone(timezone.utc)


async def _broadcast_occupancy_update(db: Session, manager) -> None:
    """Compute current occupancy stats and push to all WebSocket clients."""
    try:
        from app.models import FaceEvent, AppSetting
        from sqlalchemy import func
        from datetime import timedelta

        def _get_setting(key, default="0"):
            row = db.get(AppSetting, key)
            return row.value if row else default

        max_capacity = int(_get_setting("occupancy_max_capacity", "0") or 0)

        latest_sq = (
            db.query(
                FaceEvent.person_id,
                func.max(FaceEvent.created_at).label("max_ts"),
            )
            .filter(
                FaceEvent.source == "face_analytics",
                FaceEvent.direction.in_(["in", "out"]),
                FaceEvent.person_id.isnot(None),
            )
            .group_by(FaceEvent.person_id)
            .subquery()
        )
        currently_inside = (
            db.query(FaceEvent)
            .join(
                latest_sq,
                (FaceEvent.person_id == latest_sq.c.person_id)
                & (FaceEvent.created_at == latest_sq.c.max_ts),
            )
            .filter(
                FaceEvent.source == "face_analytics",
                FaceEvent.direction == "in",
            )
            .count()
        )

        today_start = yerevan_day_start_utc()
        today_end = today_start + timedelta(hours=24)
        in_today = db.query(FaceEvent).filter(
            FaceEvent.source == "face_analytics",
            FaceEvent.direction == "in",
            FaceEvent.created_at >= today_start,
            FaceEvent.created_at < today_end,
        ).count()
        out_today = db.query(FaceEvent).filter(
            FaceEvent.source == "face_analytics",
            FaceEvent.direction == "out",
            FaceEvent.created_at >= today_start,
            FaceEvent.created_at < today_end,
        ).count()

        pct = round(currently_inside / max_capacity * 100, 1) if max_capacity > 0 else None
        over = currently_inside > max_capacity if max_capacity > 0 else False

        await manager.broadcast({
            "type": "occupancy_update",
            "currently_inside": currently_inside,
            "max_capacity": max_capacity,
            "capacity_pct": pct,
            "is_over_capacity": over,
            "in_today": in_today,
            "out_today": out_today,
        })
    except Exception as _occ_err:
        logger.debug("Occupancy broadcast error: %s", _occ_err)


def add_audit_log(
    db: Session,
    user,
    action: str,
    target_type: Optional[str] = None,
    target_id=None,
    details: Optional[str] = None,
    ip_address: Optional[str] = None,
):
    from app.models import AuditLog
    try:
        log = AuditLog(
            user_id=user.id if user else None,
            username=user.username if user else "system",
            action=action,
            target_type=target_type,
            target_id=str(target_id) if target_id is not None else None,
            details=details,
            ip_address=ip_address,
            created_at=utc_now(),
        )
        db.add(log)
        db.commit()
    except Exception as e:
        db.rollback()
        logger.warning("Audit log write failed: %s", e)


def model_snapshot(obj, fields: list[str]) -> dict:
    return {field: getattr(obj, field, None) for field in fields}


def diff_snapshot(before: dict, after: dict, redact: set[str] | None = None) -> dict:
    redact = redact or set()
    diff = {}
    for key in sorted(set(before) | set(after)):
        if before.get(key) == after.get(key):
            continue
        old_val = "***" if key in redact and before.get(key) else before.get(key)
        new_val = "***" if key in redact and after.get(key) else after.get(key)
        diff[key] = {"before": old_val, "after": new_val}
    return diff


def public_path(path: str) -> str:
    return path.replace(BASE_DIR, "").replace("\\", "/")


def full_media_path(public_file_path: str) -> str:
    """Resolve a public-relative path to an absolute path, rejecting traversal attempts."""
    relative = public_file_path.strip("/").replace("/", os.sep)
    candidate = os.path.normpath(os.path.join(BASE_DIR, relative))
    base = os.path.normpath(BASE_DIR) + os.sep
    if not candidate.startswith(base):
        raise ValueError(f"Path traversal rejected: {public_file_path!r}")
    return candidate


def sign_media_fields(item):
    if item is None:
        return item
    fields = (
        "preview_image",
        "image_path",
        "crop_image_path",
        "original_image_path",
        "image",
        "crop_image",
        "event_image",
    )
    if isinstance(item, dict):
        for field in fields:
            if field in item:
                item[field] = _create_media_url(item[field])
        return item
    for field in fields:
        if hasattr(item, field):
            setattr(item, field, _create_media_url(getattr(item, field)))
    return item


def sign_media_collection(items):
    return [sign_media_fields(item) for item in items]


async def read_limited(file: UploadFile, max_bytes: Optional[int] = None, require_image: bool = True) -> bytes:
    content_type = (file.content_type or "").lower()
    if require_image and content_type and not content_type.startswith("image/"):
        raise HTTPException(415, "Only image uploads are allowed")
    limit = max_bytes if max_bytes is not None else MAX_UPLOAD_BYTES
    chunks = []
    total = 0
    while True:
        chunk = await file.read(1024 * 1024)
        if not chunk:
            break
        total += len(chunk)
        if total > limit:
            raise HTTPException(413, f"File too large. Maximum: {limit // (1024 * 1024)} MB")
        chunks.append(chunk)
    return b"".join(chunks)


def save_bytes_to_media(data: bytes, prefix="face") -> str:
    day = utc_now().strftime("%Y-%m-%d")
    folder = os.path.join(MEDIA_DIR, day)
    os.makedirs(folder, exist_ok=True)
    filename = f"{prefix}_{utc_now().strftime('%H%M%S_%f')}.jpg"
    full_path = os.path.join(folder, filename)
    with open(full_path, "wb") as f:
        f.write(data)
    return public_path(full_path)


def _smooth_age(db, person_id: int, new_age: int, window: int = 8) -> int:
    """
    B) Weighted average of recent age estimates for the same person.
    Newest estimate has highest weight (2^0=1 for oldest … 2^(n-1) for newest).
    Prevents outliers (bad lighting, partial face) from corrupting the display age.
    """
    try:
        from app.models import FaceEvent as _FE
        recent = (
            db.query(_FE.age)
            .filter(_FE.person_id == person_id)
            .filter(_FE.age.isnot(None))
            .filter(_FE.age > 0)
            .order_by(_FE.created_at.desc())
            .limit(window - 1)
            .all()
        )
        ages = [new_age] + [r.age for r in recent if r.age and r.age > 0]
        if len(ages) < 2:
            return new_age
        # Exponential weights: newest first → [2^(n-1), ..., 2, 1]
        n = len(ages)
        weights = [2 ** i for i in range(n - 1, -1, -1)]
        smoothed = sum(a * w for a, w in zip(ages, weights)) / sum(weights)
        return max(1, round(smoothed))
    except Exception:
        return new_age


def crop_face_preview(image_path: str, bbox, prefix="face_crop") -> Optional[str]:
    img = cv2.imread(image_path)
    if img is None:
        return None
    h, w = img.shape[:2]
    x1, y1, x2, y2 = [int(v) for v in bbox]
    face_w = max(1, x2 - x1)
    face_h = max(1, y2 - y1)
    size = int(max(face_w, face_h) * 1.9)
    cx = int((x1 + x2) / 2)
    cy = int((y1 + y2) / 2)
    nx1 = max(0, cx - size // 2)
    ny1 = max(0, cy - size // 2)
    nx2 = min(w, cx + size // 2)
    ny2 = min(h, cy + size // 2)
    crop = img[ny1:ny2, nx1:nx2]
    if crop.size == 0:
        return None
    crop = cv2.resize(crop, (320, 320))
    day = utc_now().strftime("%Y-%m-%d")
    folder = os.path.join(MEDIA_DIR, day)
    os.makedirs(folder, exist_ok=True)
    filename = f"{prefix}_{utc_now().strftime('%H%M%S_%f')}.jpg"
    full_path = os.path.join(folder, filename)
    ok = cv2.imwrite(full_path, crop)
    if not ok or not os.path.exists(full_path):
        return None
    return public_path(full_path)


def draw_face_box(image_path: str, bbox, prefix="face_box") -> Optional[str]:
    img = cv2.imread(image_path)
    if img is None:
        return None
    x1, y1, x2, y2 = [int(v) for v in bbox]
    cv2.rectangle(img, (x1, y1), (x2, y2), (0, 255, 0), 3)
    day = utc_now().strftime("%Y-%m-%d")
    folder = os.path.join(MEDIA_DIR, day)
    os.makedirs(folder, exist_ok=True)
    filename = f"{prefix}_{utc_now().strftime('%H%M%S_%f')}.jpg"
    full_path = os.path.join(folder, filename)
    ok = cv2.imwrite(full_path, img)
    if not ok or not os.path.exists(full_path):
        return None
    return public_path(full_path)


async def extract_faces_async(image_path: str):
    loop = asyncio.get_running_loop()
    return await loop.run_in_executor(_face_executor, extract_faces, image_path)


def estimate_face_quality(image_path: str, bbox, det_score: float = 0.0) -> dict:
    img = cv2.imread(image_path)
    if img is None:
        return {"ok": False, "score": 0.0, "reason": "image_read_failed"}
    h, w = img.shape[:2]
    x1, y1, x2, y2 = [int(v) for v in bbox]
    x1, y1 = max(0, x1), max(0, y1)
    x2, y2 = min(w, x2), min(h, y2)
    face_w = max(0, x2 - x1)
    face_h = max(0, y2 - y1)
    if face_w < MIN_FACE_SIZE or face_h < MIN_FACE_SIZE:
        return {"ok": False, "score": 0.0, "reason": "face_too_small",
                "det_score": float(det_score or 0.0), "blur_var": 0.0,
                "face_width": face_w, "face_height": face_h}
    crop = img[y1:y2, x1:x2]
    if crop.size == 0:
        return {"ok": False, "score": 0.0, "reason": "empty_crop"}
    gray = cv2.cvtColor(crop, cv2.COLOR_BGR2GRAY)
    blur_var = float(cv2.Laplacian(gray, cv2.CV_64F).var())
    size_score = min(100.0, (min(face_w, face_h) / 160.0) * 100.0)
    det_score_norm = max(0.0, min(100.0, float(det_score or 0.0) * 100.0))
    blur_score = max(0.0, min(100.0, (blur_var / 120.0) * 100.0))
    quality = round((det_score_norm * 0.45) + (size_score * 0.25) + (blur_score * 0.30), 2)
    ok = det_score >= MIN_FACE_DET_SCORE and quality >= MIN_FACE_QUALITY_SCORE
    if ok:
        reason = "ok"
    elif det_score < MIN_FACE_DET_SCORE:
        reason = "low_detection_confidence"
    elif blur_var < 45:
        reason = "blurry_face"
    else:
        reason = "low_quality"
    return {"ok": ok, "score": quality, "reason": reason, "det_score": float(det_score or 0.0),
            "blur_var": blur_var, "face_width": face_w, "face_height": face_h}


def get_event_type(person) -> str:
    if person.person_type == "employee":
        return "employee_detected"
    if person.person_type == "vip":
        return "vip_detected"
    if person.person_type == "blacklist":
        return "blacklist_detected"
    return "unknown_detected"


def find_registered_person(db: Session, embedding):
    from app.models import Person
    from app.faiss_engine import FAISS_AVAILABLE, search_registered_person
    if FAISS_AVAILABLE:
        try:
            person, score = search_registered_person(db, Person, embedding, MATCH_THRESHOLD)
            if person:
                return person, score
        except Exception as e:
            logger.warning("FAISS registered search failed: %s", e)

    persons = (db.query(Person).filter(Person.face_embedding.isnot(None))
               .filter(Person.person_type != "unknown").all())
    best_person, best_score = None, 0.0
    for person in persons:
        try:
            stored = json.loads(person.face_embedding)
            score = compare_embeddings(embedding, stored)
            if score > best_score:
                best_score = score
                best_person = person
        except Exception:
            continue
    if best_person and best_score >= MATCH_THRESHOLD:
        return best_person, best_score
    return None, best_score


def find_unknown_group(db: Session, embedding):
    from app.models import Person
    from app.faiss_engine import FAISS_AVAILABLE, search_unknown_group
    if FAISS_AVAILABLE:
        try:
            person, score = search_unknown_group(db, Person, embedding, UNKNOWN_GROUP_THRESHOLD)
            if person:
                return person, score
        except Exception as e:
            logger.warning("FAISS unknown search failed: %s", e)

    persons = (db.query(Person).filter(Person.face_embedding.isnot(None))
               .filter(Person.person_type == "unknown").all())
    best_person, best_score = None, 0.0
    for person in persons:
        try:
            stored = json.loads(person.face_embedding)
            score = compare_embeddings(embedding, stored)
            if score > best_score:
                best_score = score
                best_person = person
        except Exception:
            continue
    if best_person and best_score >= UNKNOWN_GROUP_THRESHOLD:
        return best_person, best_score
    return None, best_score


def export_response(filepath: str, filename: str, media_type: str, bg):
    """Return FileResponse and schedule temp file cleanup after download."""
    from fastapi.responses import FileResponse
    bg.add_task(_delete_file, filepath)
    return FileResponse(filepath, media_type=media_type, filename=filename)


def _delete_file(path: str):
    try:
        if os.path.exists(path):
            os.remove(path)
    except Exception as e:
        logger.warning("Failed to delete temp file %s: %s", path, e)


def media_file_response(media_path: str):
    """Serve a file from MEDIA_DIR, blocking path traversal."""
    from fastapi.responses import FileResponse
    relative = media_path.strip("/").replace("/", os.sep)
    if relative.startswith("media" + os.sep):
        relative = relative[len("media" + os.sep):]
    full_path = os.path.abspath(os.path.join(MEDIA_DIR, relative))
    media_root = os.path.abspath(MEDIA_DIR)
    if full_path != media_root and not full_path.startswith(media_root + os.sep):
        from fastapi import HTTPException
        raise HTTPException(400, "Invalid media path")
    if not os.path.exists(full_path) or not os.path.isfile(full_path):
        from fastapi import HTTPException
        raise HTTPException(404, "Media not found")
    return FileResponse(full_path)


def style_excel_sheet(ws):
    from openpyxl.styles import Font, Alignment
    from openpyxl.utils import get_column_letter
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            max_len = max(max_len, len(str(cell.value or "")))
        ws.column_dimensions[col_letter].width = min(max_len + 3, 45)


def event_to_payload(event, person) -> dict:
    person_name = person.full_name or getattr(person, "unknown_label", None) or f"Unknown #{person.id}"
    return sign_media_fields({
        "type": "new_event",
        "event_id": event.id,
        "person_id": person.id,
        "person_name": person_name,
        "person_type": person.person_type,
        "camera_name": event.camera_name,
        "image": event.image_path,
        "crop_image": event.crop_image_path,
        "confidence": event.confidence,
        "face_index": getattr(event, "face_index", None),
        "face_count": getattr(event, "face_count", None),
        "face_quality": getattr(event, "face_quality", None),
        "quality_reason": getattr(event, "quality_reason", None),
        "is_low_quality": getattr(event, "is_low_quality", 0),
        "direction": getattr(event, "direction", None),
        "age": getattr(event, "age", None),
        "gender": getattr(event, "gender", None),
        "created_at": event.created_at.isoformat() if event.created_at else None,
    })


async def process_faces_and_create_events(
    db: Session,
    image_full_path: str,
    original_image_path: str,
    camera_ip,
    camera_name: str,
    raw_payload,
    prefix: str,
    direction: str = "both",
    camera_id: int | None = None,
):
    from app.models import FaceEvent
    from app.ws_manager import manager

    faces = await extract_faces_async(image_full_path)
    created = []
    skipped = []

    if not faces:
        return {"events": [], "skipped": [], "face_count": 0}

    total_faces = len(faces)

    for idx, face in enumerate(faces, start=1):
        bbox = face.get("bbox")
        det_score = float(face.get("det_score") or 0.0)
        quality = estimate_face_quality(image_full_path, bbox, det_score)

        threshold = get_face_quality_threshold(db)
        is_low_quality = 0 if (det_score >= MIN_FACE_DET_SCORE and quality.get("score", 0) >= threshold) else 1

        embedding = face["embedding"]
        crop_image = crop_face_preview(image_full_path, bbox, f"{prefix}_crop_{idx}") or original_image_path
        boxed_image = draw_face_box(image_full_path, bbox, f"{prefix}_box_{idx}") or original_image_path

        person, confidence, matched_registered, matched_unknown = upsert_person_for_face(
            db, embedding, crop_image, camera_name,
        )

        face_age    = face.get("age")
        face_gender = face.get("gender")

        # B) Weighted age smoothing — average of last 8 estimates, newest = highest weight
        if face_age and person and person.id:
            face_age = _smooth_age(db, person.id, face_age)

        event = FaceEvent(
            person_id=person.id,
            camera_ip=camera_ip,
            camera_id=camera_id,
            camera_name=camera_name,
            event_type="low_quality" if is_low_quality else get_event_type(person),
            source="face_analytics",
            image_path=boxed_image,
            crop_image_path=crop_image,
            original_image_path=original_image_path,
            confidence=confidence,
            raw_payload=raw_payload,
            direction=direction,
            age=face_age,
            gender=face_gender,
            created_at=utc_now(),
        )

        for attr, value in {
            "face_index": idx,
            "face_count": total_faces,
            "face_quality": quality.get("score"),
            "quality_reason": quality.get("reason"),
            "det_score": det_score,
            "matched_unknown_group": 1 if matched_unknown else 0,
            "is_low_quality": is_low_quality,
        }.items():
            if hasattr(event, attr):
                setattr(event, attr, value)

        db.add(event)
        db.flush()
        db.refresh(event)
        db.commit()

        # Auto-update attendance record for registered persons
        try:
            from app.attendance_engine import upsert_attendance_record
            upsert_attendance_record(db, person, event.created_at)
            db.commit()
        except Exception as _ae_err:
            logger.debug("Attendance engine skipped: %s", _ae_err)

        try:
            from app.alert_engine import check_and_fire_alerts
            check_and_fire_alerts(db, event, person)
        except Exception as _ae:
            logger.debug("Alert engine error: %s", _ae)

        broadcast_payload = event_to_payload(event, person)
        if face_age:
            broadcast_payload["age"] = face_age
        if face_gender:
            broadcast_payload["gender"] = face_gender

        try:
            from app.webhook_engine import dispatch_webhooks
            dispatch_webhooks(db, event.event_type or "unknown_detected", broadcast_payload)
        except Exception as _we:
            logger.debug("Webhook engine error: %s", _we)

        try:
            from app.door_engine import fire_door_triggers
            fire_door_triggers(db, person, confidence)
        except Exception as _de:
            logger.debug("Door trigger error: %s", _de)

        try:
            from app.face_door_engine import handle_face_event
            import asyncio as _asyncio
            _asyncio.ensure_future(handle_face_event(
                db=db,
                person_id=person.id if person else None,
                camera_id=camera_id,
                confidence=confidence,
                person_type=person.person_type if person else None,
                broadcast_fn=manager.broadcast,
            ))
        except Exception as _fde:
            logger.debug("Face door engine error: %s", _fde)

        await manager.broadcast(broadcast_payload)

        if direction in ("in", "out"):
            await _broadcast_occupancy_update(db, manager)

        created.append({
            "event": event,
            "person": person,
            "confidence": confidence,
            "matched_registered": matched_registered,
            "matched_unknown_group": matched_unknown,
            "face_index": idx,
            "face_quality": quality,
        })

    return {"events": created, "skipped": skipped, "face_count": total_faces}


def upsert_person_for_face(db: Session, embedding, crop_image: str, camera_name: str):
    from app.models import Person
    from app.faiss_engine import mark_faiss_dirty
    embedding_json = json.dumps(embedding.tolist())
    registered, reg_score = find_registered_person(db, embedding)
    if registered:
        registered.last_seen = utc_now()
        registered.seen_count = (registered.seen_count or 0) + 1
        if not registered.preview_image:
            registered.preview_image = crop_image
        if hasattr(registered, "last_camera_name"):
            registered.last_camera_name = camera_name
        db.flush()
        db.refresh(registered)
        return registered, reg_score, True, False

    unknown, unk_score = find_unknown_group(db, embedding)
    if unknown:
        unknown.last_seen = utc_now()
        unknown.seen_count = (unknown.seen_count or 0) + 1
        unknown.preview_image = unknown.preview_image or crop_image
        if hasattr(unknown, "last_camera_name"):
            unknown.last_camera_name = camera_name
        db.flush()
        db.refresh(unknown)
        return unknown, unk_score, False, True

    person = Person(
        full_name=None, person_type="unknown", source="fa", preview_image=crop_image,
        face_embedding=embedding_json, first_seen=utc_now(), last_seen=utc_now(), seen_count=1,
    )
    for attr, val in [("is_unknown_group", 1), ("unknown_label", "Unknown"), ("last_camera_name", camera_name)]:
        if hasattr(person, attr):
            setattr(person, attr, val)
    db.add(person)
    db.flush()
    db.refresh(person)
    if hasattr(person, "unknown_label"):
        person.unknown_label = f"Unknown #{person.id}"
        db.flush()
        db.refresh(person)
    mark_faiss_dirty()
    return person, None, False, False
