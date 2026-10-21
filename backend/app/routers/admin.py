"""Admin router - audit logs, FAISS status, jobs, exports (Excel/PDF), media."""

import ipaddress
import logging
import os
import secrets
from typing import List, Optional

from fastapi import APIRouter, BackgroundTasks, Depends, HTTPException, Query, Request
from fastapi.responses import Response
from jose import JWTError, jwt
from sqlalchemy.orm import Session, joinedload

from app.auth import ALGORITHM, SECRET_KEY, require_admin
from app.db import get_db
from app.faiss_engine import FAISS_AVAILABLE, rebuild_faiss_indexes
from app.job_queue import get_job, list_jobs
from app.media_security import verify_media_signature
from app.metrics import render_prometheus_metrics
from app.models import AuditLog, FaceEvent, Person, User
from app.schemas import AuditLogOut
from app.utils import (
    EXPORT_DIR,
    YEREVAN_TZ,
    add_audit_log,
    export_response,
    media_file_response,
    style_excel_sheet,
    utc_now,
)

logger = logging.getLogger(__name__)
router = APIRouter(tags=["admin"])


def _user_from_raw_token(raw: Optional[str], db: Session) -> User:
    if not raw:
        raise HTTPException(401, "Not authenticated")
    try:
        payload = jwt.decode(raw, SECRET_KEY, algorithms=[ALGORITHM])
        user_id = payload.get("sub")
        if not user_id:
            raise HTTPException(401, "Invalid token")
    except JWTError:
        raise HTTPException(401, "Invalid token")
    user = db.query(User).filter(User.id == int(user_id)).first()
    if not user or not user.is_active:
        raise HTTPException(401, "Inactive or unknown user")
    return user


@router.get("/media/{media_path:path}")
def protected_media(
    media_path: str,
    request: Request,
    token: Optional[str] = Query(None),
    exp: Optional[int] = Query(None),
    sig: Optional[str] = Query(None),
    db: Session = Depends(get_db),
):
    public_path_value = f"/media/{media_path}"
    if not verify_media_signature(public_path_value, exp, sig):
        bearer = request.headers.get("Authorization", "")
        raw = token or (bearer[7:] if bearer.startswith("Bearer ") else None)
        _user_from_raw_token(raw, db)
    return media_file_response(media_path)


@router.get("/api/audit-logs", response_model=List[AuditLogOut])
def get_audit_logs(
    limit: int = 300,
    action: Optional[str] = None,
    username: Optional[str] = None,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_admin),
):
    query = db.query(AuditLog)
    if action:
        query = query.filter(AuditLog.action == action)
    if username:
        query = query.filter(AuditLog.username.ilike(f"%{username}%"))
    return query.order_by(AuditLog.created_at.desc()).limit(limit).all()


@router.get("/api/faiss/status")
def faiss_status(
    db: Session = Depends(get_db),
    current_user: User = Depends(require_admin),
):
    ok = False
    message = "FAISS unavailable"
    if FAISS_AVAILABLE:
        try:
            rebuild_faiss_indexes(db, Person)
            ok = True
            message = "FAISS enabled and indexes rebuilt"
        except Exception as e:
            message = f"FAISS rebuild failed: {e}"
    return {"faiss_available": FAISS_AVAILABLE, "ok": ok, "message": message}


@router.get("/api/jobs")
def jobs(
    limit: int = 100,
    current_user: User = Depends(require_admin),
):
    return {"jobs": list_jobs(limit=min(max(limit, 1), 500))}


@router.get("/api/jobs/{job_id}")
def job_detail(
    job_id: str,
    current_user: User = Depends(require_admin),
):
    job = get_job(job_id)
    if not job:
        raise HTTPException(404, "Job not found")
    return job


def _metrics_access_allowed(request: Request) -> bool:
    configured_key = os.getenv("METRICS_API_KEY", "").strip()
    supplied_key = request.headers.get("X-Metrics-Key", "").strip()
    if configured_key and supplied_key and secrets.compare_digest(configured_key, supplied_key):
        return True

    client_ip = request.client.host if request.client else ""
    try:
        address = ipaddress.ip_address(client_ip)
    except ValueError:
        return False
    if address.is_loopback:
        return True
    return os.getenv("METRICS_ALLOW_PRIVATE", "0") == "1" and address.is_private


@router.get("/metrics")
def metrics(request: Request):
    if not _metrics_access_allowed(request):
        raise HTTPException(403, "Metrics access denied")
    return Response(render_prometheus_metrics(), media_type="text/plain; version=0.0.4")


# -- Exports ------------------------------------------------------------------

@router.get("/api/export/attendance")
def export_attendance_excel(
    bg: BackgroundTasks,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_admin),
):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"
    ws.append(["ID", "Person", "Type", "Camera", "Created At", "Confidence", "Quality"])

    events = (
        db.query(FaceEvent)
        .options(joinedload(FaceEvent.person))
        .filter(FaceEvent.source == "face_analytics")
        .order_by(FaceEvent.created_at.desc())
        .limit(5000)
        .all()
    )

    for e in events:
        person_name = "Unknown Person"
        person_type = "unknown"
        if e.person:
            person_name = e.person.full_name or getattr(e.person, "unknown_label", None) or f"Unknown #{e.person.id}"
            person_type = e.person.person_type or "unknown"
        ws.append([e.id, person_name, person_type, e.camera_name, str(e.created_at), e.confidence, getattr(e, "face_quality", None)])

    style_excel_sheet(ws)
    filename = f"attendance_export_{utc_now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    filepath = os.path.join(EXPORT_DIR, filename)
    wb.save(filepath)

    add_audit_log(db=db, user=current_user, action="EXPORT_EXCEL", target_type="attendance",
                  details=f"Exported attendance Excel: {filename}")
    return export_response(filepath, filename, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", bg)


@router.get("/api/export/events")
def export_events_excel(
    bg: BackgroundTasks,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_admin),
):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Events"
    ws.append(["ID", "Person", "Person Type", "Event Type", "Camera", "Camera IP",
               "Confidence", "Quality", "Low Quality", "Created At"])

    events = (
        db.query(FaceEvent)
        .options(joinedload(FaceEvent.person))
        .filter(FaceEvent.source == "face_analytics")
        .order_by(FaceEvent.created_at.desc())
        .limit(5000)
        .all()
    )

    for e in events:
        person_name = "Unknown Person"
        person_type = "unknown"
        if e.person:
            person_name = e.person.full_name or getattr(e.person, "unknown_label", None) or f"Unknown #{e.person.id}"
            person_type = e.person.person_type or "unknown"
        ws.append([e.id, person_name, person_type, e.event_type, e.camera_name, e.camera_ip,
                   e.confidence, getattr(e, "face_quality", None), getattr(e, "is_low_quality", 0), str(e.created_at)])

    style_excel_sheet(ws)
    filename = f"events_export_{utc_now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    filepath = os.path.join(EXPORT_DIR, filename)
    wb.save(filepath)

    add_audit_log(db=db, user=current_user, action="EXPORT_EXCEL", target_type="events",
                  details=f"Exported events Excel: {filename}")
    return export_response(filepath, filename, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", bg)


@router.get("/api/export/persons")
def export_persons_excel(
    bg: BackgroundTasks,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_admin),
):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Persons"
    ws.append(["ID", "Full Name", "Type", "Department", "Phone", "Seen Count",
               "First Seen", "Last Seen", "Last Camera", "Notes"])

    persons = (
        db.query(Person)
        .filter(Person.full_name.isnot(None))
        .filter(Person.person_type != "unknown")
        .filter(Person.source != "acs")
        .order_by(Person.last_seen.desc())
        .limit(5000)
        .all()
    )

    for p in persons:
        ws.append([p.id, p.full_name, p.person_type, p.department, p.phone, p.seen_count,
                   str(p.first_seen), str(p.last_seen), getattr(p, "last_camera_name", None), p.notes])

    style_excel_sheet(ws)
    filename = f"persons_export_{utc_now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    filepath = os.path.join(EXPORT_DIR, filename)
    wb.save(filepath)

    add_audit_log(db=db, user=current_user, action="EXPORT_EXCEL", target_type="persons",
                  details=f"Exported persons Excel: {filename}")
    return export_response(filepath, filename, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", bg)


@router.get("/api/export/audit-logs")
def export_audit_logs_excel(
    bg: BackgroundTasks,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_admin),
):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Audit Logs"
    ws.append(["ID", "Username", "Action", "Target Type", "Target ID", "Details", "IP Address", "Created At"])

    logs = db.query(AuditLog).order_by(AuditLog.created_at.desc()).limit(5000).all()

    for log in logs:
        ws.append([log.id, log.username, log.action, log.target_type, log.target_id,
                   log.details, log.ip_address, str(log.created_at)])

    style_excel_sheet(ws)
    filename = f"audit_logs_export_{utc_now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    filepath = os.path.join(EXPORT_DIR, filename)
    wb.save(filepath)

    add_audit_log(db=db, user=current_user, action="EXPORT_EXCEL", target_type="audit_logs",
                  details=f"Exported audit logs Excel: {filename}")
    return export_response(filepath, filename, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", bg)


@router.get("/api/export/attendance/pdf")
def export_attendance_pdf(
    bg: BackgroundTasks,
    limit: int = 1000,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_admin),
):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    filename = f"attendance_report_{utc_now().strftime('%Y%m%d_%H%M%S')}.pdf"
    filepath = os.path.join(EXPORT_DIR, filename)

    doc = SimpleDocTemplate(filepath, pagesize=A4, rightMargin=24, leftMargin=24, topMargin=28, bottomMargin=24)
    styles = getSampleStyleSheet()
    elements = [
        Paragraph("<b>GSC LLC - Face Analytics Pro</b>", styles["Title"]),
        Paragraph("Attendance Report", styles["Heading2"]),
        Paragraph(
            f"Generated at: {utc_now().astimezone(YEREVAN_TZ).strftime('%Y-%m-%d %H:%M:%S')} Asia/Yerevan",
            styles["Normal"],
        ),
        Spacer(1, 16),
    ]

    events = (
        db.query(FaceEvent)
        .options(joinedload(FaceEvent.person))
        .filter(FaceEvent.source == "face_analytics")
        .order_by(FaceEvent.created_at.desc())
        .limit(min(limit, 5000))
        .all()
    )

    data = [["Time", "Person", "Type", "Camera", "Confidence", "Quality"]]
    for e in events:
        person_name = "Unknown Person"
        person_type = "unknown"
        if e.person:
            person_name = e.person.full_name or getattr(e.person, "unknown_label", None) or f"Unknown #{e.person.id}"
            person_type = e.person.person_type or "unknown"

        confidence = f"{round(float(e.confidence) * 100, 1)}%" if e.confidence is not None else "-"
        face_quality = getattr(e, "face_quality", None)
        quality = str(round(float(face_quality), 1)) if face_quality is not None else "-"

        created = "-"
        if e.created_at:
            try:
                created = e.created_at.astimezone(YEREVAN_TZ).strftime("%Y-%m-%d %H:%M:%S")
            except Exception:
                created = str(e.created_at)

        data.append([created, person_name, person_type, e.camera_name or "-", confidence, quality])

    table = Table(data, repeatRows=1, colWidths=[95, 110, 70, 120, 65, 55])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f2937")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#cbd5e1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))

    elements.append(table)
    doc.build(elements)

    add_audit_log(db=db, user=current_user, action="EXPORT_PDF", target_type="attendance",
                  details=f"Exported attendance PDF report: {filename}")
    return export_response(filepath, filename, "application/pdf", bg)
